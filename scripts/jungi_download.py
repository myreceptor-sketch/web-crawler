"""전기넷(jungi.net) 공고별 '참여업체 엑셀' 다운로드 스크립트 (사용자 PC에서 실행).

공고마다: 상세페이지 열기 → 공고번호/종목(전기) 확인 → 참여업체보기 클릭
→ 업체목록 창에서 엑셀 클릭 → 다운로드 폴더에 저장.
엑셀은 참여업체 전체를 담는다. 다만 업체목록이 다 불러와지기 전에 누르면 50행만 받아지므로,
충분히 기다린 뒤 누르고 정확히 50행이면 더 기다려 한 번 더 받는다. 파일명은 <공고번호>.xlsx.

사용:
    pip install playwright openpyxl && playwright install chromium
    python scripts/jungi_download.py --xlsx "<목록.xlsx>" --limit 3        # 먼저 3건 시험
    python scripts/jungi_download.py --xlsx "<목록.xlsx>"                   # 전체 (끊겨도 재실행하면 이어서)
    python scripts/jungi_download.py --xlsx "<목록.xlsx>" --only E012502290  # 한 공고만 시험

기본 저장 폴더는 사용자 Downloads(C:\\Users\\admin\\Downloads). 진행 기록은 --log (CSV).
로그인이 필요하면 창이 뜬 상태(기본 headed)에서 직접 로그인한 뒤 Enter 를 누른다.
ID/PW 는 코드·파일에 저장하지 않는다.
"""
import argparse
import csv
import random
import re
import sys
import time
from pathlib import Path

import openpyxl
from playwright.sync_api import TimeoutError as PWTimeout
from playwright.sync_api import sync_playwright

HEADER_ROW = 6          # 엑셀 6행이 헤더, 7행부터 데이터
COL_NO, COL_URL = 3, 10  # 0-based: 공고번호(D), 상세페이지(K)
MAX_CONSECUTIVE_FAIL = 5
SEARCH_URL = "https://www.jungi.net/nbid"   # 낙찰정보 검색 (공사명·공고번호·발주기관)
CAP_ROWS = 50            # 목록이 덜 불러와졌을 때 엑셀에 담기는 행 수 (의심 기준)
LOAD_WAIT = 5.0          # 업체목록 창을 연 뒤 엑셀을 누르기 전 기다리는 시간(초)


def cell_url(v):
    """셀 값에서 URL 을 꺼낸다. =HYPERLINK("https://…","표시글") 수식이면 첫 번째 주소를 쓴다."""
    m = re.search(r"https?://[^\"\s]+", str(v))
    return m.group(0) if m else str(v).strip()


def read_targets(xlsx, col_no=COL_NO, col_url=COL_URL, header_row=HEADER_ROW):
    """col_url < 0 이면 URL 열이 없는 목록 — 공고번호만 읽고 URL 은 빈 문자열(검색 모드)."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    out = []
    for row in wb.active.iter_rows(min_row=header_row + 1, values_only=True):
        if col_url < 0:
            if len(row) > col_no and row[col_no] and re.fullmatch(r"E\d+(-\d+)?", str(row[col_no]).strip()):
                out.append((str(row[col_no]).strip(), ""))
        elif len(row) > max(col_no, col_url) and row[col_no] and row[col_url]:
            out.append((str(row[col_no]).strip(), cell_url(row[col_url])))
    wb.close()
    return out


def read_expected(xlsx, col_no, col_exp):
    """목록 엑셀의 '전체 참여업체 수' 열 → {공고번호: 수}. 받은 파일의 업체 수 검증에 쓴다."""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    out = {}
    for row in wb.active.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if len(row) > max(col_no, col_exp) and row[col_no] and str(row[col_exp]).strip().isdigit():
            out[str(row[col_no]).strip()] = int(str(row[col_exp]).strip())
    wb.close()
    return out


def load_done(log):
    """완료로 본 공고번호. 업체 수 검증 도입 전 기록(상세가 'rows=' 로 시작하지 않음)은 제외한다."""
    if not log.exists():
        return set()
    with open(log, encoding="utf-8-sig") as f:
        return {r["공고번호"] for r in csv.DictReader(f)
                if r["상태"] == "OK" and r["상세"].startswith("rows=")}


def safe_name(no):
    return re.sub(r'[\\/:*?"<>|]', "_", no)


def result_file(outdir, no):
    return outdir / f"{safe_name(no)}.xlsx"


def count_rows(path):
    """받은 엑셀의 업체 행 수 (순위 열이 숫자인 행). 읽지 못하면 -1."""
    wb = None
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        # 위쪽 공고 정보 블록은 제외 — 순위(숫자, 낙찰하한가 미달은 음수)와 사업자등록번호(10자리)가 함께 있는 행만 센다
        return sum(1 for r in wb.active.iter_rows(values_only=True)
                   if len(r) > 1 and r[0] is not None and re.fullmatch(r"-?\d+", str(r[0]).strip())
                   and re.fullmatch(r"\d{10}", re.sub(r"\D", "", str(r[1] or ""))))
    except Exception:
        return -1
    finally:
        if wb:
            wb.close()   # read_only 로 열면 파일이 계속 잠겨 삭제·덮어쓰기가 막힌다


def unique_path(folder, name):
    """브라우저처럼 같은 이름이 있으면 'name (1).ext' 로 저장한다 (덮어쓰지 않음)."""
    p = folder / name
    n = 1
    while p.exists():
        p = folder / f"{Path(name).stem} ({n}){Path(name).suffix}"
        n += 1
    return p


def find_click(ctx_pages, text, timeout=8000):
    """열려 있는 모든 페이지/프레임에서 text 가 보이는 요소를 찾아 반환."""
    end = time.time() + timeout / 1000
    while time.time() < end:
        for pg in ctx_pages():
            for fr in pg.frames:
                for loc in (fr.get_by_role("button", name=re.compile(text)),
                            fr.get_by_role("link", name=re.compile(text)),
                            fr.get_by_text(text, exact=False)):
                    try:
                        if loc.count() and loc.first.is_visible():
                            return loc.first
                    except Exception:
                        pass
        time.sleep(0.3)
    return None


def download_excel(ctx, pop, dest):
    """업체목록 창의 엑셀 버튼을 눌러 dest 로 저장. 실패하면 사유 문자열, 성공이면 None."""
    xl = find_click(lambda: [pop], r"엑셀|Excel")
    if not xl:
        return "엑셀 버튼 없음"
    # 엑셀 버튼이 새 탭을 열어 다운로드하는 경우도 있으므로, 이후 생기는 페이지에도 리스너를 건다
    got = []

    def attach(pg):
        pg.on("download", lambda d: got.append(d))

    for p in ctx.pages:
        attach(p)
    ctx.on("page", attach)
    try:
        xl.click()
        end = time.time() + 60
        while not got and time.time() < end:
            pop.wait_for_timeout(300)  # time.sleep 은 Playwright 이벤트를 처리하지 못한다
    finally:
        ctx.remove_listener("page", attach)
    if not got:
        return "다운로드 이벤트 없음"
    got[0].save_as(dest)
    return None


def open_by_search(ctx, page, no):
    """상세 URL 이 없는 목록용: 낙찰정보(/nbid) 검색창에 공고번호를 넣어 검색하고, 결과 행을 눌러 상세로 간다.

    반환: (상세를 보여주는 페이지, None) 또는 (None, 실패 사유). 검색창이 없으면 로그인이 풀린 것이다.
    """
    page.goto(SEARCH_URL, wait_until="domcontentloaded", timeout=30000)
    page.wait_for_load_state("networkidle", timeout=15000)
    box = page.locator("#search_text")
    if not box.count():
        return None, "검색창 없음 — 로그인이 풀렸는지 확인"
    box.fill(no)
    page.locator("#search_btn").click()
    page.wait_for_load_state("networkidle", timeout=15000)
    page.wait_for_timeout(1000)
    row = page.locator("tr", has_text=no)
    if not row.count():
        return None, "검색 결과 없음"
    before = set(ctx.pages)
    link = row.first.locator("a")
    (link.first if link.count() else row.first).click()
    page.wait_for_timeout(2000)
    opened = list(set(ctx.pages) - before)
    detail = opened[0] if opened else page    # 새 탭이면 그 탭, 아니면 같은 탭
    try:
        detail.wait_for_load_state("networkidle", timeout=15000)
    except PWTimeout:
        pass
    return detail, None


def process(ctx, page, no, url, outdir, delay=8.0, force=False, expected=None, search=False):
    if search:
        page, err = open_by_search(ctx, page, no)
        if err:
            return "SKIP" if err == "검색 결과 없음" else "FAIL", err
    else:
        page.goto(url, wait_until="domcontentloaded", timeout=30000)
        page.wait_for_load_state("networkidle", timeout=15000)
    body = page.inner_text("body")

    if "/login" in page.url:   # 세션이 풀리면 로그인 화면으로 보내진다 — SKIP 으로 넘기지 말고 연속 실패로 멈춘다
        return "FAIL", "로그인이 풀림 — 창에서 다시 로그인 필요"
    if no not in body:
        return "SKIP", "공고번호 불일치"
    m = re.search(r"종목\s*[:：]?\s*([^\n\t]+)", body)
    kind = m.group(1).strip() if m else ""
    if "전기" not in kind:
        return "SKIP", f"종목={kind or '확인불가'}"

    before = set(ctx.pages)
    btn = find_click(lambda: ctx.pages, "참여업체보기")
    if not btn:
        return "FAIL", "참여업체보기 버튼 없음"
    # '참여업체보기' 는 새 창(all_company)을 연다. 그 창이 열린 것을 확인한 뒤에만 그 창의 엑셀을 누른다.
    # (상세 페이지에도 엑셀 버튼이 있지만 그건 50행만 담긴다 — 절대 누르지 않는다)
    try:
        with ctx.expect_page(timeout=15000) as ev:
            btn.click()
        pop = ev.value
        pop.wait_for_load_state("domcontentloaded", timeout=15000)
    except PWTimeout:
        return "FAIL", "참여업체 보기 창이 열리지 않음"
    if "all_company" not in pop.url and "참여업체 보기" not in pop.inner_text("body"):
        return "FAIL", f"참여업체 보기 창이 아님: {pop.url[:80]}"

    dest = result_file(outdir, no)
    if dest.exists() and (force or count_rows(dest) == CAP_ROWS):
        try:
            dest.unlink()       # 상세 페이지 엑셀(50행 이하)로 잘못 받은 파일은 지우고 다시 받는다
        except PermissionError:
            return "FAIL", f"{dest.name} 이(가) 엑셀 등에서 열려 있음 — 닫고 다시 실행"
    note = ""
    if not dest.exists():
        for attempt, wait in enumerate((LOAD_WAIT, LOAD_WAIT * 3)):
            try:
                pop.wait_for_load_state("networkidle", timeout=15000)
            except PWTimeout:
                pass
            pop.wait_for_timeout(int(wait * 1000))
            err = download_excel(ctx, pop, dest)
            if err:
                return "FAIL", err
            if count_rows(dest) != CAP_ROWS:
                break
            if attempt == 0:
                dest.unlink()   # 정확히 50행이면 덜 불러온 것일 수 있어 더 기다려 다시 받는다
            else:
                note = " (정확히 50행 — 실제 업체 수 확인 필요)"
    n = count_rows(dest)
    if expected and n != expected:
        note += f" (목록 {expected} != 실제 {n} — 확인 필요)"

    for p in set(ctx.pages) - before:   # 업체목록 창 닫기
        try:
            p.close()
        except Exception:
            pass
    return "OK", f"rows={n}{note}"


def safe_accept(d):
    """alert/confirm 을 닫는다. 이미 닫힌 대화상자면 오류가 나므로 무시한다."""
    try:
        d.accept()
    except Exception:
        pass


def main():
    for st in (sys.stdout, sys.stderr):
        try:
            st.reconfigure(encoding="utf-8", errors="replace")   # 파일로 리다이렉트해도 한글·기호 출력이 깨지지 않게
        except Exception:
            pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--header-row", type=int, default=HEADER_ROW, help="헤더 행 번호 (1-based, 기본 6 — 데이터는 그 다음 행부터)")
    ap.add_argument("--col-no", type=int, default=COL_NO, help="공고번호 열 (0-based, 기본 3=D)")
    ap.add_argument("--col-url", type=int, default=COL_URL, help="상세페이지 URL 열 (0-based, 기본 10=K). URL 열이 없으면 -1 — 공고번호로 검색해서 연다")
    ap.add_argument("--col-expected", type=int, default=-1, help="전체 참여업체 수 열 (0-based, 받은 파일 검증용)")
    ap.add_argument("--force", action="store_true", help="이미 있는 파일도 지우고 다시 받는다")
    ap.add_argument("--out", default=str(Path.home() / "Downloads"))
    ap.add_argument("--log", default="jungi_download_log.csv")
    ap.add_argument("--start", type=int, default=0, help="목록에서 시작할 0-based 위치")
    ap.add_argument("--limit", type=int, default=0, help="이번 실행에서 처리할 최대 건수 (0=전체)")
    ap.add_argument("--delay", type=float, default=8.0, help="공고 사이 기본 대기(초), ±50%% 지터")
    ap.add_argument("--batch", type=int, default=50, help="이 건수마다 길게 쉰다 (0=쉬지 않음)")
    ap.add_argument("--batch-pause", type=float, default=120.0, help="배치마다 쉬는 시간(초)")
    ap.add_argument("--only", default="", help="이 공고번호 한 건만 처리 (시험용)")
    ap.add_argument("--headless", action="store_true")
    ap.add_argument("--profile", default=".jungi_profile", help="로그인 유지용 전용 브라우저 프로필 폴더")
    a = ap.parse_args()

    outdir = Path(a.out)
    outdir.mkdir(parents=True, exist_ok=True)
    log = Path(a.log)
    done = load_done(log)
    all_t = read_targets(a.xlsx, a.col_no, a.col_url, a.header_row)[a.start:]
    if a.only:
        all_t = [t for t in all_t if t[0] == a.only]
        done.discard(a.only)
    # 이전에 받은 건: 50행 미만이면 전부 받은 것, 정확히 50행이면 덜 받았을 수 있어 다시 받는다
    expected = read_expected(a.xlsx, a.col_no, a.col_expected) if a.col_expected >= 0 else {}
    for no, _ in all_t:
        f = result_file(outdir, no)
        if no not in done and f.exists() and not a.only and not a.force and 0 <= count_rows(f) < CAP_ROWS:
            done.add(no)
    targets = [t for t in all_t if t[0] not in done]
    if a.limit:
        targets = targets[: a.limit]
    print(f"대상 {len(targets)}건 (이미 완료 {len(done)}건 제외) → {outdir}")

    new_log = not log.exists()
    fails = 0
    t0 = time.time()
    with sync_playwright() as pw, open(log, "a", newline="", encoding="utf-8-sig") as lf:
        w = csv.writer(lf)
        if new_log:
            w.writerow(["공고번호", "상태", "상세", "시각"])
        ctx = pw.chromium.launch_persistent_context(
            a.profile, headless=a.headless, accept_downloads=True,
            viewport={"width": 1280, "height": 900})
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.on("dialog", safe_accept)
        page.goto("https://www.jungi.net/")
        if not a.headless:
            input("로그인이 필요하면 창에서 직접 로그인한 뒤 Enter 를 누르세요 (불필요하면 바로 Enter): ")

        for i, (no, url) in enumerate(targets, 1):
            try:
                status, detail = process(ctx, page, no, url, outdir, a.delay, a.force, expected.get(no), search=a.col_url < 0)
            except Exception as e:
                status, detail = "FAIL", f"{type(e).__name__}: {str(e)[:120]}"
            # 남은 팝업 정리
            for p in ctx.pages:
                if p is not page:
                    try:
                        p.close()
                    except Exception:
                        pass
            w.writerow([no, status, detail, time.strftime("%Y-%m-%d %H:%M:%S")])
            lf.flush()
            fails = fails + 1 if status == "FAIL" else 0
            el = time.time() - t0
            eta = el / i * (len(targets) - i)
            print(f"[{i}/{len(targets)}] {no} {status} {detail}  (남은 예상 {eta/60:.0f}분)")
            if fails >= MAX_CONSECUTIVE_FAIL:
                print(f"연속 {fails}건 실패 — 중단합니다. 사이트 구조/차단 여부를 확인하세요.")
                break
            time.sleep(a.delay * random.uniform(0.5, 1.5))
            if a.batch and i % a.batch == 0 and i < len(targets):
                print(f"{i}건마다 휴식 — {a.batch_pause:.0f}초 대기 (서버 부하 완화)")
                time.sleep(a.batch_pause)
        ctx.close()


if __name__ == "__main__":
    sys.exit(main())
